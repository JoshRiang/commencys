# Build the historical chapter Markdown as editable Word documents.
# Images are resolved relative to each chapter file so the source bundle remains portable.
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.opc.constants import RELATIONSHIP_TYPE as RT


# Chapter sources and generated DOCX files live together in the frozen history bundle.
ROOT = Path(__file__).resolve().parents[1]
RESULT = ROOT / "HISTORY_TRACE" / "RESULT"
PAGE_WIDTH_IN = 8.27
SIDE_MARGIN_IN = 0.92
CONTENT_WIDTH_IN = PAGE_WIDTH_IN - 2 * SIDE_MARGIN_IN
INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
IMAGE_PATTERN = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")


def set_run_font(run, name="Times New Roman", size=12):
    # Set both common Word font slots so Latin and East Asian text render consistently.
    run.font.name = name
    run.font.size = Pt(size)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def add_hyperlink(paragraph, label, url):
    # Create an external hyperlink with the document's body font and link styling.
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run_node = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    font_node = OxmlElement("w:rFonts")
    font_node.set(qn("w:ascii"), "Times New Roman")
    font_node.set(qn("w:hAnsi"), "Times New Roman")
    properties.append(font_node)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F5D8F")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    run_node.append(properties)
    text_node = OxmlElement("w:t")
    text_node.text = label
    run_node.append(text_node)
    hyperlink.append(run_node)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text):
    # Preserve the chapter's bold, italic, code, and link markup while creating Word runs.
    cursor = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            set_run_font(run)
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            set_run_font(run)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, "Consolas", 10)
        else:
            label, url = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run)


def set_cell_shading(cell, fill):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "5")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D9E0E7")


def add_table(document, rows):
    column_count = len(rows[0])
    table = document.add_table(rows=len(rows), cols=column_count)
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table_width = CONTENT_WIDTH_IN
    weights = []
    for column in range(column_count):
        longest = max(len(row[column]) for row in rows)
        weights.append(min(max(longest, 8), 56))
    total = sum(weights)
    widths = [table_width * weight / total for weight in weights]
    floor = table_width * (0.12 if column_count > 2 else 0.18)
    widths = [max(width, floor) for width in widths]
    scale = table_width / sum(widths)
    widths = [width * scale for width in widths]

    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            row._tr.get_or_add_trPr().append(header)
        for column, value in enumerate(values):
            cell = row.cells[column]
            cell.width = Inches(widths[column])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.line_spacing = 1.12
            paragraph.paragraph_format.space_after = Pt(1)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline(paragraph, value.strip())
            for run in paragraph.runs:
                set_run_font(run, size=9.5)
                if row_index == 0:
                    run.bold = True
            if row_index == 0:
                set_cell_shading(cell, "EAF1F7")
    set_table_borders(table)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.space_before = Pt(0)


def split_table_row(line):
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def configure_document(document, title):
    set_section_layout(document.sections[0], landscape=False)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.first_line_indent = Inches(0.28)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for style_name, size, before, after in (
        ("Title", 18, 0, 14),
        ("Heading 1", 15, 16, 8),
        ("Heading 2", 13, 12, 6),
    ):
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = None
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    document.core_properties.title = title


def set_section_layout(section, landscape):
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11.69)
        section.page_height = Inches(8.27)
        section.top_margin = Inches(0.62)
        section.bottom_margin = Inches(0.62)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(PAGE_WIDTH_IN)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.88)
        section.bottom_margin = Inches(0.88)
        section.left_margin = Inches(SIDE_MARGIN_IN)
        section.right_margin = Inches(SIDE_MARGIN_IN)


def build_one(markdown_path):
    source = markdown_path.read_text(encoding="utf-8-sig").splitlines()
    document = Document()
    title = next((line.lstrip("# ").strip() for line in source if line.startswith("# ")), markdown_path.stem)
    configure_document(document, title)
    paragraph_index = 0
    index = 0
    in_fence = False
    while index < len(source):
        line = source[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("~~~") or line.startswith("```"):
            in_fence = not in_fence
            index += 1
            continue
        if in_fence:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.first_line_indent = Inches(0)
            paragraph.paragraph_format.line_spacing = 1.1
            paragraph.paragraph_format.space_after = Pt(3)
            run = paragraph.add_run(line)
            set_run_font(run, "Consolas", 9)
            index += 1
            continue
        image_match = IMAGE_PATTERN.fullmatch(line)
        if image_match:
            _, image_reference = image_match.groups()
            image_path = (markdown_path.parent / image_reference).resolve()
            if not image_path.is_file():
                raise FileNotFoundError(f"Image not found for {markdown_path.name}: {image_path}")
            is_diagram = "DIAGRAMS" in image_reference or "UML_PICTURE" in image_reference
            if is_diagram:
                section = document.add_section(WD_SECTION_START.NEW_PAGE)
                set_section_layout(section, landscape=True)
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Inches(0)
            paragraph.paragraph_format.space_before = Pt(5)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.keep_with_next = True
            figure_width = 10.2 if is_diagram else CONTENT_WIDTH_IN
            paragraph.add_run().add_picture(str(image_path), width=Inches(figure_width))
            caption_index = index + 1
            while caption_index < len(source) and not source[caption_index].strip():
                caption_index += 1
            if is_diagram and caption_index < len(source):
                caption = source[caption_index].strip()
                if caption.startswith("*") and caption.endswith("*"):
                    caption_paragraph = document.add_paragraph(style="Caption")
                    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    caption_paragraph.paragraph_format.first_line_indent = Inches(0)
                    add_inline(caption_paragraph, caption)
                    for run in caption_paragraph.runs:
                        set_run_font(run, size=10)
                    index = caption_index
            if is_diagram:
                section = document.add_section(WD_SECTION_START.NEW_PAGE)
                set_section_layout(section, landscape=False)
            index += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while index < len(source) and source[index].strip().startswith("|"):
                candidate = source[index].strip()
                if not re.fullmatch(r"\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?", candidate):
                    table_lines.append(split_table_row(candidate))
                index += 1
            if table_lines:
                add_table(document, table_lines)
            continue
        if line.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Inches(0)
            add_inline(paragraph, line[2:].strip())
            index += 1
            continue
        if line.startswith("## ") or line.startswith("### "):
            level = 1 if line.startswith("## ") else 2
            paragraph = document.add_paragraph(style=f"Heading {level}")
            paragraph.paragraph_format.first_line_indent = Inches(0)
            add_inline(paragraph, line[level + 1:].strip())
            index += 1
            continue
        if re.match(r"^[-*]\s+", line) or re.match(r"^\d+\.\s+", line):
            numbered = bool(re.match(r"^\d+\.\s+", line))
            content = re.sub(r"^(?:[-*]|\d+)\.??\s+", "", line)
            paragraph = document.add_paragraph(style="List Number" if numbered else "List Bullet")
            paragraph.paragraph_format.first_line_indent = Inches(0)
            paragraph.paragraph_format.left_indent = Inches(0.3)
            paragraph.paragraph_format.line_spacing = 1.25
            paragraph.paragraph_format.space_after = Pt(4)
            add_inline(paragraph, content)
            index += 1
            continue
        if line.startswith("*") and line.endswith("*"):
            paragraph = document.add_paragraph(style="Caption")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Inches(0)
            paragraph.paragraph_format.space_before = Pt(1)
            paragraph.paragraph_format.space_after = Pt(12)
            add_inline(paragraph, line)
            paragraph_index += 1
            index += 1
            continue
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.keep_together = True
        add_inline(paragraph, line)
        paragraph_index += 1
        index += 1
    output_path = RESULT / f"{markdown_path.stem}.docx"
    document.save(output_path)
    return output_path


if __name__ == "__main__":
    for chapter in (1, 2, 3):
        path = build_one(RESULT / f"CHAPTER_{chapter}.md")
        print(path)
